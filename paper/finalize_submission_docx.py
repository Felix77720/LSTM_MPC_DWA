"""Apply final submission style checks to the generated manuscript DOCX.

The manuscript remains sourced from EI_paper_final_CN.md; this pass enforces
the requested all-black text and neutral table styling after conversion.
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


HERE = Path(__file__).resolve().parent
INPUT = HERE / "EI_paper_final_CN.docx"
OUTPUT = HERE / "EI_paper_submission.docx"


def set_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_white_fill(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), "FFFFFF")


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            yield p
        for p in sec.footer.paragraphs:
            yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def main():
    doc = Document(INPUT)
    for p in all_paragraphs(doc):
        for run in p.runs:
            set_black(run)
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                set_white_fill(cell)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_black(run)
    doc.core_properties.title = "LSTM-Enhanced MPC-DWA for Dynamic Obstacle Avoidance of UAVs"
    doc.core_properties.author = "付贞辉"
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(f"WROTE {OUTPUT}")


if __name__ == "__main__":
    main()
