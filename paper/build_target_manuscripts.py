"""Build venue-specific manuscripts while preserving the original final paper.

The sources of truth remain EI_paper_final_CN.md and jirs_manuscript_EN.md.  The
two DOCX files are deliberately generated into target-specific directories so
that later venue edits cannot overwrite the preserved manuscript.
"""
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
RESULTS = HERE.parent / "results"
CN_SOURCE = HERE / "EI_paper_final_CN.md"
EN_SOURCE = HERE / "jirs_manuscript_EN.md"
SYS_TEMPLATE = HERE / "target_sys_ele" / "sys_ele_template.docx"
SYS_OUT_DIR = HERE / "target_sys_ele"
JIRS_OUT_DIR = HERE / "jirs"

sys.path.insert(0, str(HERE))
from build_final_paper_cn import (  # noqa: E402
    clean_inline,
    configure_section,
    latex_display_to_text,
    parse_markdown,
    set_cell_shading,
    set_cell_text,
)


def set_run_font(run, ascii_font="Times New Roman", east_font="SimSun", size=10.0,
                 bold=False, italic=False, color=(0, 0, 0)):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(doc, *, chinese=False, template=False):
    sec = doc.sections[0]
    if not template:
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.top_margin = Mm(18)
        sec.bottom_margin = Mm(18)
        sec.left_margin = Mm(16.5 if chinese else 25.4)
        sec.right_margin = Mm(16.5 if chinese else 25.4)
    else:
        configure_section(sec, 1)
    normal = doc.styles["Normal"]
    set_run_font(normal, ascii_font="Times New Roman", east_font="SimSun" if chinese else "SimSun", size=9.5 if chinese else 10.5)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(3)
    for name, size in (("Heading 1", 12.5), ("Heading 2", 10.5), ("Heading 3", 9.5)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_run_font(style, ascii_font="Arial", east_font="SimSun" if chinese else "Times New Roman", size=size, bold=True)
        style.paragraph_format.space_before = Pt(7 if name == "Heading 1" else 5)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True
    for section in doc.sections:
        section.header.paragraphs[0].text = ""
        add_page_number(section.footer.paragraphs[0])


def add_text_paragraph(doc, text, *, chinese=False, kind="body", bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    elif kind in {"title", "meta", "label", "equation", "caption"}:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind in {"title", "meta", "equation"} else WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0 if kind in {"title", "meta", "label", "caption"} else 1.04
    p.paragraph_format.first_line_indent = Pt(0 if kind != "body" else (14 if chinese else 0))
    p.paragraph_format.space_before = Pt(4 if kind in {"label", "caption"} else 0)
    p.paragraph_format.space_after = Pt(2 if kind != "caption" else 4)
    if kind == "caption":
        p.paragraph_format.keep_with_next = True
    run = p.add_run(clean_inline(text))
    size = {"title": 16 if chinese else 16, "meta": 10.5, "label": 10, "caption": 9, "body": 9.5 if chinese else 10.5}.get(kind, 10)
    set_run_font(run, ascii_font="Arial" if kind == "title" else "Times New Roman",
                 east_font="SimSun" if chinese else "Times New Roman",
                 size=size, bold=bold, italic=italic)
    return p


def add_equation(doc, text, *, chinese=False):
    # Convert display-math commands to readable Unicode before inserting the
    # result into the editable Word math container.
    text = latex_display_to_text(text)
    text = re.sub(r"hat\{([A-Za-z]+)\}", lambda m: "".join(ch + "̂" for ch in m.group(1)), text)
    text = text.replace("||", "‖")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.set(qn("xml:space"), "preserve")
    math_text.text = text
    math_run.append(math_text)
    omath.append(math_run)
    omath_para.append(omath)
    p._p.append(omath_para)
    return p


def add_table(doc, rows, *, chinese=False):
    ncols = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    base_size = 7.7 if chinese else 8.1
    for j in range(ncols):
        value = rows[0][j] if j < len(rows[0]) else ""
        set_cell_text(table.rows[0].cells[j], clean_inline(value), True, base_size, (0, 0, 0))
        set_cell_shading(table.rows[0].cells[j], "FFFFFF")
    for row in rows[1:]:
        cells = table.add_row().cells
        for j in range(ncols):
            value = row[j] if j < len(row) else ""
            set_cell_text(cells[j], clean_inline(value), False, base_size - 0.2, (0, 0, 0))
            set_cell_shading(cells[j], "FFFFFF")
    width_in = 6.98 if chinese else 6.3
    for idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if idx == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cell in row.cells:
            cell.width = Inches(width_in / ncols)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, source_path, caption, *, chinese=False):
    source_path = source_path.resolve()
    if not source_path.exists():
        print(f"WARNING missing figure: {source_path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(source_path), width=Inches(6.75 if chinese else 6.25))
    cap = add_text_paragraph(doc, caption, chinese=chinese, kind="caption", align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.keep_with_next = False
    if cap.runs:
        set_run_font(cap.runs[0], ascii_font="Times New Roman", east_font="SimSun" if chinese else "Times New Roman", size=8.8, italic=False)


def resolve_figure(source, rel_path):
    return (source.parent / rel_path.replace("%20", " ")).resolve()


def body_blocks(source, start_heading):
    blocks = parse_markdown(source)
    started = False
    selected = []
    for block in blocks:
        if block[0] == "h2" and block[1].strip() == start_heading:
            started = True
        if started:
            selected.append(block)
    return selected


def add_blocks(doc, blocks, *, source, chinese=False):
    ordered_index = 0
    for index, (kind, value) in enumerate(blocks):
        if kind == "h1":
            add_text_paragraph(doc, value, chinese=chinese, kind="title", bold=True)
        elif kind == "h2":
            p = doc.add_heading(clean_inline(value), level=1)
            p.paragraph_format.keep_with_next = True
        elif kind == "h3":
            p = doc.add_heading(clean_inline(value), level=2)
            p.paragraph_format.keep_with_next = True
        elif kind == "equation":
            add_equation(doc, value, chinese=chinese)
        elif kind == "table":
            add_table(doc, value, chinese=chinese)
        elif kind == "figure":
            caption, rel = value
            add_figure(doc, resolve_figure(source, rel), caption, chinese=chinese)
        elif kind in {"bullet", "number"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if kind == "number":
                ordered_index += 1
                prefix = f"{ordered_index}. "
            else:
                prefix = "• "
            run = p.add_run(prefix + clean_inline(value))
            set_run_font(run, east_font="SimSun" if chinese else "Times New Roman", size=9.5 if chinese else 10.5)
        else:
            text = clean_inline(value)
            if text.startswith(("表 ", "Table ", "图 ", "Fig. ")):
                add_text_paragraph(doc, text, chinese=chinese, kind="caption", bold=True)
            elif text.startswith("本文数值来自") or text.startswith("All numeric"):
                add_text_paragraph(doc, text, chinese=chinese, kind="caption")
            else:
                add_text_paragraph(doc, text, chinese=chinese, kind="body")


def add_sys_ele_body(doc, blocks, source):
    """Use the official Chinese-journal rhythm: text in two columns and
    equations/tables/figures temporarily spanning the page width."""
    ordered_index = 0
    columns = 2
    for index, (kind, value) in enumerate(blocks):
        next_kind = blocks[index + 1][0] if index + 1 < len(blocks) else None
        text_value = clean_inline(value) if kind == "paragraph" else ""
        wide = kind in {"table", "figure", "equation"} or text_value.startswith("表 ")
        if wide and columns != 1:
            sec = doc.add_section(WD_SECTION.CONTINUOUS)
            configure_section(sec, 1)
            columns = 1
        elif not wide and columns != 2:
            sec = doc.add_section(WD_SECTION.CONTINUOUS)
            configure_section(sec, 2)
            columns = 2
        if kind == "h2":
            p = doc.add_heading(clean_inline(value), level=1)
            p.paragraph_format.keep_with_next = True
        elif kind == "h3":
            p = doc.add_heading(clean_inline(value), level=2)
            p.paragraph_format.keep_with_next = True
        elif kind == "equation":
            add_equation(doc, value, chinese=True)
        elif kind == "table":
            add_table(doc, value, chinese=True)
        elif kind == "figure":
            caption, rel = value
            add_figure(doc, resolve_figure(source, rel), caption, chinese=True)
        elif kind in {"bullet", "number"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if kind == "number":
                ordered_index += 1
                prefix = f"{ordered_index}. "
            else:
                prefix = "• "
            run = p.add_run(prefix + clean_inline(value))
            set_run_font(run, east_font="SimSun", size=9.5)
        else:
            if text_value.startswith(("表 ", "图 ")):
                add_text_paragraph(doc, text_value, chinese=True, kind="caption", bold=True)
            elif text_value.startswith("本文数值来自"):
                add_text_paragraph(doc, text_value, chinese=True, kind="caption")
            else:
                add_text_paragraph(doc, text_value, chinese=True, kind="body")


def set_core(doc, title, author):
    props = doc.core_properties
    props.title = title
    props.author = author
    props.subject = "Venue-specific manuscript"
    props.keywords = "UAV, dynamic obstacle avoidance, LSTM, MPC, DWA"
    props.comments = ""


def build_sys_ele():
    SYS_OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SYS_TEMPLATE)) if SYS_TEMPLATE.exists() else Document()
    clear_body(doc)
    configure_document(doc, chinese=True, template=SYS_TEMPLATE.exists())
    add_text_paragraph(doc, "LSTM增强MPC-DWA无人机避障方法", chinese=True, kind="title", bold=True)
    add_text_paragraph(doc, "付贞辉", chinese=True, kind="meta")
    add_text_paragraph(doc, "单位：", chinese=True, kind="meta")
    add_text_paragraph(doc, "邮编：", chinese=True, kind="meta")
    add_text_paragraph(doc, "LSTM-Enhanced MPC-DWA for UAV Obstacle Avoidance", chinese=False, kind="meta", bold=True)
    add_text_paragraph(doc, "FU Zhenhui", chinese=False, kind="meta")
    add_text_paragraph(doc, "Affiliation: To be completed by the author", chinese=False, kind="meta")
    add_text_paragraph(doc, "摘要", chinese=True, kind="label", bold=True)
    cn_abstract = ("针对三维动态障碍环境中观测、预测与规划不同步导致的避障失效，提出将LSTM轨迹预测同步注入模型预测控制（MPC）和动态窗口法（DWA）的耦合方法。以最近1 s历史预测未来3 s位置，并将同一预测用于MPC安全代价和DWA净空评价。基于800条独立轨迹进行轨迹级训练/验证，采用训练集归一化、噪声延迟增强、早停和最佳模型选择。固定场景LSTM成功12/12，12个随机场景成功36/36。结果表明，该方法在所测试仿真条件下改善了预测—规划一致性，但联合感知扰动下仍存在失败，尚不能证明普适鲁棒性。")
    add_text_paragraph(doc, cn_abstract, chinese=True, kind="body")
    add_text_paragraph(doc, "关键词：无人机；动态避障；模型预测控制；动态窗口法；LSTM；Kalman滤波；轨迹预测", chinese=True, kind="body")
    add_text_paragraph(doc, "中图分类号：TP273", chinese=True, kind="body")
    add_text_paragraph(doc, "Abstract", chinese=False, kind="label", bold=True)
    en_front = ("Dynamic obstacles can invalidate local UAV planning when observation, prediction and control are not temporally aligned. This paper introduces an LSTM trajectory predictor into a bidirectionally coupled model predictive control (MPC) and dynamic window approach (DWA) framework. The predictor receives the latest 1 s of an obstacle history and forecasts its three-dimensional position over the next 3 s; the same time-indexed forecast is injected into the MPC safety cost and the DWA candidate-clearance score. The revised protocol uses 800 independent trajectories, a trajectory-level 640/160 split, training-only normalization statistics, correctly aligned noisy/delayed-history augmentation, early stopping and best-validation-loss selection. LSTM succeeds in 12/12 fixed trials and 36/36 random trials under the tested clean simulation suite. The evidence supports improved prediction–planning consistency under the tested conditions, but does not establish universal robustness to arbitrary sensor delay or real-flight conditions.")
    add_text_paragraph(doc, en_front, chinese=False, kind="body")
    add_text_paragraph(doc, "Keywords: UAV; dynamic obstacle avoidance; model predictive control; dynamic window approach; LSTM; Kalman filter; trajectory prediction", chinese=False, kind="body")
    blocks = body_blocks(CN_SOURCE, "1 引言")
    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(body_section, 2)
    add_sys_ele_body(doc, blocks, source=CN_SOURCE)
    author_section = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(author_section, 1)
    add_text_paragraph(doc, "作者简介", chinese=True, kind="label", bold=True)
    add_text_paragraph(doc, "付贞辉（学历、职称待补），研究方向为无人机动态避障、模型预测控制和轨迹预测。E-mail：待补。", chinese=True, kind="body")
    for section in doc.sections[1:]:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
    set_core(doc, "LSTM增强MPC-DWA无人机避障方法", "付贞辉")
    out = SYS_OUT_DIR / "system_engineering_electronics_manuscript_CN.docx"
    doc.save(str(out))
    print(f"WROTE {out}")


def build_jirs():
    JIRS_OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc, chinese=False)
    add_text_paragraph(doc, "Prediction–Planning Consistent LSTM-Enhanced MPC-DWA for UAV Dynamic Obstacle Avoidance", chinese=False, kind="title", bold=True)
    add_text_paragraph(doc, "Zhenhui Fu", chinese=False, kind="meta")
    add_text_paragraph(doc, "Affiliation: To be completed by the author", chinese=False, kind="meta")
    add_text_paragraph(doc, "Corresponding author: To be completed by the author", chinese=False, kind="meta")
    add_text_paragraph(doc, "Category: (2) Systems Modelling/Simulation/Control/Computer-Aided Design; (5) Intelligent Systems/Intelligent Control/Robot Motion Planning", chinese=False, kind="meta")
    add_text_paragraph(doc, "Abstract", chinese=False, kind="label", bold=True)
    en_blocks = parse_markdown(EN_SOURCE)
    in_body = False
    body = []
    for block in en_blocks:
        if block[0] == "h2" and block[1].strip() == "1 Introduction":
            in_body = True
        if in_body:
            body.append(block)
    abstract = next((clean_inline(value) for kind, value in en_blocks if kind == "paragraph" and value.startswith("Dynamic obstacles can invalidate")), "")
    add_text_paragraph(doc, abstract, chinese=False, kind="body")
    add_text_paragraph(doc, "Keywords: UAV; dynamic obstacle avoidance; model predictive control; dynamic window approach; LSTM; Kalman filter; trajectory prediction", chinese=False, kind="body")
    add_blocks(doc, body, source=EN_SOURCE, chinese=False)
    set_core(doc, "Prediction–Planning Consistent LSTM-Enhanced MPC-DWA for UAV Dynamic Obstacle Avoidance", "Zhenhui Fu")
    out = JIRS_OUT_DIR / "jirs_manuscript_EN.docx"
    doc.save(str(out))
    print(f"WROTE {out}")


if __name__ == "__main__":
    build_sys_ele()
    build_jirs()
