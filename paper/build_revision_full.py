"""Single-source revision builder.  Legacy builder is imported unchanged."""
from pathlib import Path
import argparse, importlib.util, re
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent; RESULTS=HERE.parent/'results'
OUT=HERE/'EI_paper_revision_CN.docx'; TEX=HERE/'EI_paper_revision_CN.tex'
FIXED=RESULTS/'prediction_fixed'/'fixed_window_summary.csv'
LABEL={'static':'静态冻结','cv':'匀速外推','lstm':'本文 LSTM'}
def build_content(fixed_df, loop_df=None, protocol=None):
    rows=[]
    for k in ('static','cv','lstm'):
        z=fixed_df[fixed_df.Method.astype(str).str.lower()==k]
        rows.append((LABEL[k],len(z),z.ADE.mean(),z.FDE.mean()))
    zh=('针对三维动态障碍环境中的无人机避障，本文将长短期记忆网络（LSTM）嵌入双向耦合的MPC-DWA规划器。网络使用最近10步、步长0.1 s的障碍历史预测未来30步（3 s），预测位置同时进入MPC时变软碰撞代价和DWA时序净空评分。在闭环对照中统一采用三维路线恢复权重wRoute=3，并以wRoute=0作消融；固定窗口预测本身不涉及该规划权重。固定S1–S4离线评估覆盖17个障碍、每障碍20个重叠窗口，每法340条：静态冻结ADE/FDE %.3f/%.3f m，匀速外推%.3f/%.3f m，LSTM %.3f/%.3f m。窗口重叠不构成独立重复；预测优势不自动等于闭环安全优势。本文实证范围限于固定窗口预测评估，闭环安全和模块贡献尚未得到验证。')%(rows[0][2],rows[0][3],rows[1][2],rows[1][3],rows[2][2],rows[2][3])
    en=('We study three-dimensional UAV avoidance of dynamic obstacles with an LSTM module coupled to MPC and DWA. The network uses 10 history steps at 0.1 s intervals to predict 30 future steps (3 s); predictions enter the MPC time-varying soft collision cost and the DWA temporal-clearance score. In the closed-loop comparison, all groups use a 3D route-recovery weight of wRoute=3, with wRoute=0 as an ablation; this planning weight is not part of the fixed-window prediction test. The fixed S1–S4 offline evaluation contains 17 obstacles and 20 overlapping windows per obstacle (340 windows per method): static ADE/FDE %.3f/%.3f m, constant-velocity %.3f/%.3f m, and LSTM %.3f/%.3f m. Overlapping windows are not independent replicates, and prediction accuracy does not imply closed-loop safety. The empirical scope is limited to fixed-window prediction; closed-loop safety and module contributions remain unvalidated.')%(rows[0][2],rows[0][3],rows[1][2],rows[1][3],rows[2][2],rows[2][3])
    loop_rows=[]
    if loop_df is not None:
        for i,(_,r) in enumerate(loop_df.iterrows(),1):
            def v(k,default='—'):
                x=r.get(k,default)
                try:
                    if pd.isna(x) or x==float('inf'): return default
                except Exception: pass
                return x
            loop_rows.append((f'E{i}',str(v('Scenario')).split()[0],str(v('Method')),v('Horizon'),v('RouteWeight'),v('Success'),v('DetourEvents'),v('CollisionEvents'),v('MinClearance'),v('PathLength'),v('TrackRMSE'),v('MeanStepTime')))
    return {'abstract_zh':zh,'abstract_en':en,'table_rows':rows,'loop_rows':loop_rows}

def load_old():
    s=importlib.util.spec_from_file_location('old',HERE/'build_paper_docx_cn.py'); m=importlib.util.module_from_spec(s); s.loader.exec_module(m); m.OUT=str(OUT); return m
def load_fixed():
    d=pd.read_csv(FIXED); req={'Scenario','Obstacle','HistoryStep','Method','ADE','FDE'}
    if not req.issubset(d.columns): raise ValueError('fixed schema missing columns')
    if set(d.Method.astype(str).str.lower()) != set(LABEL): raise ValueError('fixed methods mismatch')
    counts=d.Method.astype(str).str.lower().value_counts().to_dict()
    if any(counts.get(k)!=340 for k in LABEL): raise ValueError(f'fixed counts {counts}')
    return d
def load_loop(run_dir):
    if run_dir is None: return None
    p=Path(run_dir).resolve()/'metrics.csv'
    if not p.exists(): raise FileNotFoundError(p)
    return pd.read_csv(p)
def pred(run_dir=None):
    p=FIXED
    return load_fixed() if p.exists() else None
def replace_para(p, reps):
    t=p.text
    for a,b in reps.items(): t=t.replace(a,b)
    if t!=p.text: p.text=t
def remove_range(doc):
    body=doc._element.body; active=False; doomed=[]; ref=None
    for el in list(body):
        txt=''.join(el.xpath('.//w:t/text()')) if el.tag.endswith('}p') else ''
        if txt.strip()=='IV. 实验': active=True
        if active and txt.strip()=='参考文献': ref=el; active=False; continue
        if active: doomed.append(el)
    for el in doomed: body.remove(el)
    return ref
def add_section_before(doc, ref, title, text):
    p=doc.add_heading(title,level=1); p.paragraph_format.keep_with_next=True
    q=doc.add_paragraph(text); q.paragraph_format.first_line_indent=Pt(20)
    if ref is not None:
        body=doc._element.body
        for el in (p._p,q._p): body.remove(el); body.insert(body.index(ref),el)
def build_doc(run_dir=None):
    m=load_old(); m.build_document(); doc=Document(OUT)
    reps={'结果表明本文方法在避障成功率和轨迹预测精度两方面均优于基线。':'结果分别报告预测误差与闭环安全，不预设两类指标均占优。','且 ADE/FDE 显著低于静态基线、波动最小。':'现有结果不支持显著性或波动最小推断。','本文 LSTM 保持全程成功':'旧闭环记录不作为新版结论。','本文 LSTM 在全部场景保持成功':'旧闭环记录不作为新版结论。','结果表明本文方法在避障成功率和轨迹预测精度两方面均优于基线':'结果分别报告预测误差与闭环安全，不预设两类指标均占优','while the proposed LSTM method keeps full success and far lower ADE/FDE than static freezing.':'the fixed-window experiment reports descriptive ADE/FDE; closed-loop claims await the revised protocol.'}
    for p in doc.paragraphs: replace_para(p,reps)
    ref=remove_range(doc); d=pred(run_dir); loop=load_loop(run_dir); content=build_content(d,loop)
    for p in doc.paragraphs:
        if p.text.strip().startswith('摘要——'):
            p.text='摘要——'+content['abstract_zh']
        elif p.text.strip().startswith('Abstract——'):
            p.text='Abstract——'+content['abstract_en']
    add_section_before(doc,ref,'IV. 实验','实验在 MATLAB R2025a、CPU 和 CasADi/IPOPT 环境中进行，步长为 0.1 s，MPC 预测长度为 30 步。修订协议统一使用 X(:,k+1)、trust ref(k+1)，滚动参考移位并按动力学延长，初始 pitch=atan2(dz,norm(dxy))。新增重复遭遇场景由三个独立、固定的有界动态障碍轨迹组成；无人机状态不参与障碍轨迹生成。MPC 增加三维回归原线路软代价 wRoute，并保留 wRoute=0 对照；DWA 对带符号预测净空不大于零的候选予以拒绝，若安全候选为空，仅在可达窗口中选择最小重叠候选并记录 fallback 次数。')
    if d is not None and {'Method','ADE','FDE'}.issubset(d.columns):
        add_section_before(doc,ref,'IV-A. 固定窗口离线预测','固定窗口结果存在时间重叠，记录数不等同于独立重复；以下为描述性均值。')
        t=doc.add_table(rows=1,cols=4); t.style='Table Grid'
        for c,x in zip(t.rows[0].cells,['方法','窗口数','ADE(m)','FDE(m)']): c.text=x
        for k in ('static','cv','lstm'):
            z=d[d.Method.astype(str).str.lower()==k]; cells=t.add_row().cells
            cells[0].text=LABEL[k]; cells[1].text=str(len(z)); cells[2].text=f'{z.ADE.mean():.3f}'; cells[3].text=f'{z.FDE.mean():.3f}'
        body=doc._element.body; el=t._element
        if ref is not None: body.remove(el); body.insert(body.index(ref),el)
    if loop is not None:
        add_section_before(doc,ref,'IV-B. 闭环计算代价','E1…组ID与效果表一致；均值和失败行均保留。时间单位为毫秒，失败运行不删除。')
        t2=doc.add_table(rows=1,cols=7); t2.style='Table Grid'
        for c,x in zip(t2.rows[0].cells,['ID','最小净空','路径长','TrackRMSE','均值ms','P95ms','失败']): c.text=x
        for i,r in enumerate(loop.to_dict('records'),1):
            cells=t2.add_row().cells
            def ms(k):
                x=r.get(k,'—')
                try:
                    if pd.isna(x): return '—'
                    return f'{float(x)*1000:.1f}'
                except Exception: return '—'
            vals=[f'E{i}',r.get('MinClearance','—'),r.get('PathLength','—'),r.get('TrackRMSE','—'),ms('MeanStepTime'),ms('P95StepTime'),str(1-int(r.get('Success',0)))]
            for c,x in zip(cells,vals): c.text=str(x)
        body=doc._element.body; el=t2._element
        if ref is not None: body.remove(el); body.insert(body.index(ref),el)
    loop_note = ('闭环数据未提供，闭环安全和模块贡献尚未得到验证。' if loop is None else f'截至读取时闭环 CSV 含 {len(loop)} 行；失败行保留，按协议分组汇总，不宣称完整多种子。')
    add_section_before(doc,ref,'IV-C. 闭环验证边界',loop_note+' 重复遭遇事件、无障碍对照、MPC-only/DWA-only/两层注入和 H=10/20/30 消融须按明确 run 目录分组；碰撞时间步与独立遭遇分别报告。')
    if loop is not None:
        lr=content['loop_rows']; add_section_before(doc,ref,'IV-D. 闭环结果','表中保留 CSV 全部行，失败运行不删除；N0 净空以—表示。该 quick 数据仅为接口测试，按单种子记录，不作泛化统计。')
        t=doc.add_table(rows=1,cols=9); t.style='Table Grid'
        for c,x in zip(t.rows[0].cells,['ID','场景','方法','H','wR','到达','偏移事件','接触事件','最小净空']): c.text=x
        for r in lr:
            cells=t.add_row().cells
            for c,x in zip(cells,r[:9]): c.text=str(x)
        body=doc._element.body; el=t._element
        if ref is not None: body.remove(el); body.insert(body.index(ref),el)
    add_section_before(doc,ref,'V. 结论','正式证据限于固定窗口离线预测误差；由于窗口重叠，不作独立重复或显著性推断。旧闭环和 quick 运行不支持新版闭环结论，模块贡献须待新协议数据。')
    doc.paragraphs[0].style=doc.styles['Title']; doc.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('w:pBdr'))
    doc.save(OUT)
    assert sum(p.text.strip()=='IV. 实验' for p in doc.paragraphs)==1
    alltxt='\n'.join(p.text for p in doc.paragraphs)
    assert '保持全程成功' not in alltxt and '波动最小' not in alltxt
    if loop is not None: assert len(content['loop_rows']) == len(loop)
    assert any('X(:,k+1)' in p.text for p in doc.paragraphs)
    if d is not None: assert len([x for x in doc.tables if len(x.rows)>=4])>=1

def build_tex(run_dir=None):
    text=(HERE/'EI_paper_draft_CN.tex').read_text(encoding='utf-8'); bs=chr(92); d=pred(run_dir); content=build_content(d,load_loop(run_dir))
    text=text.replace('结果表明本文方法在避障成功率和轨迹预测精度两方面均优于基线。','结果分别报告预测误差与闭环安全，不预设两类指标均占优。')
    abstract=bs+'begin{abstract}\n'+content['abstract_zh']+'\n'+bs+'end{abstract}'
    text=re.sub(r'(?s)\\begin\{abstract\}.*?\\end\{abstract\}',lambda _m: abstract,text,count=1)
    spans=list(re.finditer(r'(?s)\\begin\{abstract\}.*?\\end\{abstract\}',text))
    if spans:
        q=spans[-1]; text=text[:q.start()]+bs+'begin{abstract}\n'+content['abstract_en']+'\n'+bs+'end{abstract}'+text[q.end():]
    rows=''
    if d is not None:
        for k in ('static','cv','lstm'):
            z=d[d.Method.astype(str).str.lower()==k]; rows += '%s & %d & %.3f & %.3f%s\n'%(LABEL[k],len(z),z.ADE.mean(),z.FDE.mean(),bs*2)
    loop_rows=''
    if run_dir:
        ld=load_loop(run_dir)
        for i,r in enumerate(ld.to_dict('records'),1):
            def q(k,default='—'):
                x=r.get(k,default)
                try:
                    if pd.isna(x) or x==float('inf') or x==float('-inf'): return default
                except Exception: pass
                return x
            loop_rows += 'E%d & %s & %s & %s & %s & %s & %s & %s & %s ' % (i,str(q('Scenario')).split()[0],q('Method'),q('Horizon'),q('RouteWeight'),q('Success'),q('DetourEvents'),q('CollisionEvents'),q('MinClearance')) + bs*2 + '\n'
    exp=bs+'section{实验}\n'+bs+'subsection{时间对齐与版本边界}\n修订协议统一使用 X(:,k+1)、trust ref(k+1)，滚动参考移位并按动力学延长，初始 pitch=atan2(dz,norm(dxy))。\n'+bs+'subsection{固定窗口离线预测}\n固定评估包含 17 个障碍、每障碍 20 个窗口，即每种方法 340 条记录；窗口有时间重叠，不等于独立重复。\n'+bs+'begin{table}[t]'+bs+'centering'+bs+'caption{固定窗口预测误差}'+bs+'begin{tabular}{lrrr}'+bs+'toprule 方法 & 窗口数 & ADE(m) & FDE(m)'+bs*2+bs+'midrule\n'+rows+bs+'bottomrule'+bs+'end{tabular}'+bs+'end{table}\n'
    if loop_rows: exp += bs+'begin{table*}[t]'+bs+'centering'+bs+'caption{闭环效果表；E组ID对应计算代价表，失败行保留}'+bs+'begin{tabular}{lrrrrrrrr}'+bs+'toprule ID & 场景 & 方法 & H & w\\_R & 到达 & 偏移 & 接触 & 净空'+bs*2+bs+'midrule\n'+loop_rows+bs+'bottomrule'+bs+'end{tabular}'+bs+'end{table*}\n'
    if loop_rows:
        cost_rows=''
        for i,r in enumerate(load_loop(run_dir).to_dict('records'),1):
            def cv(k,default='—'):
                x=r.get(k,default)
                try:
                    if pd.isna(x) or x==float('inf'): return default
                except Exception: pass
                return x
            cost_rows += 'E%d & %s & %s & %s & %.1f & %.1f & %s ' % (i,cv('MinClearance'),cv('PathLength'),cv('TrackRMSE'),float(cv('MeanStepTime',0))*1000,float(cv('P95StepTime',0))*1000,cv('Success')) + bs*2 + '\n'
        exp += bs+'begin{table*}[t]'+bs+'centering'+bs+'caption{闭环计算代价；E组ID与效果表一致，时间单位为毫秒}'+bs+'begin{tabular}{lrrrrrr}'+bs+'toprule ID & 最小净空 & 路径长 & TrackRMSE & 均值(ms) & P95(ms) & 到达'+bs*2+bs+'midrule\n'+cost_rows+bs+'bottomrule'+bs+'end{tabular}'+bs+'end{table*}\n'
    exp += bs+'subsection{结论边界}\n闭环结果按逐行记录报告；重复遭遇、无障碍对照、注入层级及 H=10/20/30 的统计范围受 run 数据约束。\n'
    s=text.index(bs+'section{实验}'); c=text.index(bs+'section{结论}',s); text=text[:s]+exp+text[c:]
    s=text.index(bs+'section{结论}'); b=text.find(bs+'begin{thebibliography}',s); text=text[:s]+bs+'section{结论}\n正式证据限于固定窗口预测误差；重叠窗口不支持独立重复或显著性推断，旧闭环和 quick 运行不作为新版结论。\n\n'+text[b:]
    TEX.write_text(text,encoding='utf-8'); assert text.count(bs+'section{实验}')==1 and '340' in text and 'X(:,k+1)' in text
if __name__=='__main__':
    a=argparse.ArgumentParser(); a.add_argument('--run-dir'); ns=a.parse_args(); build_doc(ns.run_dir); build_tex(ns.run_dir); print('ASSERTIONS_OK',OUT,TEX)
